# -*- coding: utf-8 -*-
"""Generate a professional English final coop academic report (Word)."""

from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

BASE = Path(r"c:\Users\monee\OneDrive\Desktop\طلابي\ايلاف الكوير")
ASSETS = BASE / "_report_assets"
OUT = BASE / "Final_Coop_Report_Nadee_Smart_Assistant.docx"


def ensure_assets():
    ASSETS.mkdir(exist_ok=True)


def draw_box(ax, x, y, w, h, text, fc="#1A2B68", ec="#0F1C45", tc="white", fs=9):
    box = FancyBboxPatch(
        (x, y), w, h, boxstyle="round,pad=0.02,rounding_size=0.08",
        linewidth=1.2, facecolor=fc, edgecolor=ec,
    )
    ax.add_patch(box)
    ax.text(x + w / 2, y + h / 2, text, ha="center", va="center",
            color=tc, fontsize=fs, fontweight="bold", wrap=True)


def arrow(ax, x1, y1, x2, y2):
    ax.add_patch(FancyArrowPatch(
        (x1, y1), (x2, y2), arrowstyle="-|>", mutation_scale=12,
        linewidth=1.2, color="#334155",
    ))


def make_architecture_png():
    path = ASSETS / "architecture.png"
    fig, ax = plt.subplots(figsize=(10, 6.2))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 6.5)
    ax.axis("off")
    ax.set_title("System Architecture — Nadee Smart Assistant", fontsize=13, fontweight="bold", color="#1A2B68")

    draw_box(ax, 0.4, 5.1, 2.2, 0.9, "Employee\n(Web Browser)", fc="#E8F1FF", tc="#1A2B68")
    draw_box(ax, 3.5, 5.1, 3.0, 0.9, "Presentation Layer\nHTML / CSS / JS (RTL UI)", fc="#2B6CB0")
    draw_box(ax, 7.2, 5.1, 2.4, 0.9, "Auth Module\nLogin / Session", fc="#2BB673")

    draw_box(ax, 1.5, 3.4, 3.2, 0.95, "Django Views / URLs\nControllers", fc="#1A2B68")
    draw_box(ax, 5.3, 3.4, 3.5, 0.95, "Business Logic\nChat Engine + Analytics Services", fc="#203068")

    draw_box(ax, 1.2, 1.5, 3.5, 1.0, "Knowledge Base\nArticles / FAQ / Guides", fc="#F5A623", tc="#1A1A1A")
    draw_box(ax, 5.3, 1.5, 3.5, 1.0, "SQLite Database\nUsers, Chats, Domains, Metrics", fc="#F5A623", tc="#1A1A1A")

    draw_box(ax, 3.2, 0.25, 3.6, 0.8, "Official Source\nSDAIA Nadee Index Guide", fc="#EEF2FF", tc="#1A2B68")

    arrow(ax, 2.6, 5.55, 3.5, 5.55)
    arrow(ax, 6.5, 5.55, 7.2, 5.55)
    arrow(ax, 5.0, 5.1, 4.2, 4.35)
    arrow(ax, 5.0, 5.1, 6.5, 4.35)
    arrow(ax, 3.1, 3.4, 2.9, 2.5)
    arrow(ax, 6.8, 3.4, 6.8, 2.5)
    arrow(ax, 5.0, 1.5, 5.0, 1.05)

    fig.tight_layout()
    fig.savefig(path, dpi=180, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def make_sitemap_png():
    path = ASSETS / "sitemap.png"
    fig, ax = plt.subplots(figsize=(10, 5.8))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 6)
    ax.axis("off")
    ax.set_title("Application Sitemap", fontsize=13, fontweight="bold", color="#1A2B68")

    draw_box(ax, 3.5, 5.0, 3.0, 0.7, "Home Dashboard", fc="#1A2B68")
    nodes = [
        (0.3, 3.3, "Login /\nRegister"),
        (2.3, 3.3, "Smart\nAssistant"),
        (4.3, 3.3, "Knowledge\nBase"),
        (6.3, 3.3, "Analytics\nDashboard"),
        (8.3, 3.3, "About\nPlatform"),
    ]
    for x, y, t in nodes:
        draw_box(ax, x, y, 1.5, 0.95, t, fc="#2B6CB0", fs=8)
        arrow(ax, 5.0, 5.0, x + 0.75, y + 0.95)

    draw_box(ax, 1.8, 1.3, 2.2, 0.9, "Chat Session\nQ&A Flow", fc="#2BB673")
    draw_box(ax, 4.2, 1.3, 2.2, 0.9, "Article Detail\nGuide / FAQ", fc="#2BB673")
    draw_box(ax, 6.6, 1.3, 2.4, 0.9, "Domains Table\nAlerts / Forecast", fc="#2BB673")
    arrow(ax, 3.05, 3.3, 2.9, 2.2)
    arrow(ax, 5.05, 3.3, 5.3, 2.2)
    arrow(ax, 7.05, 3.3, 7.7, 2.2)

    fig.tight_layout()
    fig.savefig(path, dpi=180, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def make_usecase_png():
    path = ASSETS / "usecase.png"
    fig, ax = plt.subplots(figsize=(10, 6.5))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 7)
    ax.axis("off")
    ax.set_title("Use Case Diagram (Conceptual)", fontsize=13, fontweight="bold", color="#1A2B68")

    # actors
    ax.text(0.8, 5.8, "Employee", ha="center", fontsize=10, fontweight="bold", color="#1A2B68")
    ax.plot([0.8], [5.2], "o", markersize=18, color="#1A2B68")
    ax.plot([0.8, 0.8], [5.05, 4.4], color="#1A2B68", lw=2)
    ax.plot([0.5, 1.1], [4.85, 4.85], color="#1A2B68", lw=2)
    ax.plot([0.8, 0.55], [4.4, 3.9], color="#1A2B68", lw=2)
    ax.plot([0.8, 1.05], [4.4, 3.9], color="#1A2B68", lw=2)

    ax.text(9.2, 5.8, "Admin", ha="center", fontsize=10, fontweight="bold", color="#1A2B68")
    ax.plot([9.2], [5.2], "o", markersize=18, color="#1A2B68")
    ax.plot([9.2, 9.2], [5.05, 4.4], color="#1A2B68", lw=2)
    ax.plot([8.9, 9.5], [4.85, 4.85], color="#1A2B68", lw=2)
    ax.plot([9.2, 8.95], [4.4, 3.9], color="#1A2B68", lw=2)
    ax.plot([9.2, 9.45], [4.4, 3.9], color="#1A2B68", lw=2)

    # system boundary
    boundary = FancyBboxPatch(
        (2.0, 0.6), 6.0, 5.8, boxstyle="round,pad=0.02,rounding_size=0.1",
        linewidth=1.5, facecolor="#F8FAFC", edgecolor="#1A2B68", linestyle="--",
    )
    ax.add_patch(boundary)
    ax.text(5.0, 6.15, "Nadee Smart Assistant", ha="center", fontsize=10, fontweight="bold", color="#1A2B68")

    cases = [
        (3.3, 5.0, "Login to System"),
        (3.3, 4.0, "Ask Chatbot"),
        (3.3, 3.0, "Browse Knowledge"),
        (3.3, 2.0, "View Analytics"),
        (6.2, 4.5, "Manage Articles"),
        (6.2, 3.2, "Update Metrics"),
        (6.2, 2.0, "Manage Users"),
    ]
    for x, y, t in cases:
        draw_box(ax, x, y, 2.0, 0.7, t, fc="#FFFFFF", ec="#2B6CB0", tc="#1A2B68", fs=8)

    for y in [5.35, 4.35, 3.35, 2.35]:
        ax.plot([1.2, 3.3], [4.7, y], color="#64748B", lw=1)
    for y in [4.85, 3.55, 2.35]:
        ax.plot([8.9, 8.2], [4.7, y], color="#64748B", lw=1)

    fig.tight_layout()
    fig.savefig(path, dpi=180, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def make_flow_png():
    path = ASSETS / "chat_flow.png"
    fig, ax = plt.subplots(figsize=(10, 3.8))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 3)
    ax.axis("off")
    ax.set_title("Chatbot Answer Flow", fontsize=13, fontweight="bold", color="#1A2B68")
    steps = [
        (0.3, "User\nQuestion"),
        (2.3, "Normalize\n& Tokenize"),
        (4.3, "Match\nKnowledge"),
        (6.3, "Rank Best\nArticle"),
        (8.3, "Return Answer\n+ Source"),
    ]
    for i, (x, t) in enumerate(steps):
        draw_box(ax, x, 1.1, 1.6, 1.0, t, fc="#2B6CB0" if i % 2 == 0 else "#1A2B68", fs=8)
        if i < len(steps) - 1:
            arrow(ax, x + 1.6, 1.6, steps[i + 1][0], 1.6)
    fig.tight_layout()
    fig.savefig(path, dpi=180, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def set_run_font(run, size=12, bold=False, italic=False, color=None, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def add_para(doc, text, size=12, bold=False, italic=False, center=False, space_after=8, first_indent=True, color=None):
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if first_indent and not center:
        pf.first_line_indent = Cm(0.75)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return p


def add_h(doc, text, level=1):
    # Use real heading styles for TOC compatibility
    style = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3"}[level]
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in p.runs:
        set_run_font(run, size={1: 16, 2: 14, 3: 12}[level], bold=True, color=RGBColor(0x1A, 0x2B, 0x68), name="Calibri")
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    for run in p.runs:
        set_run_font(run, size=12)
    return p


def shade_cell(cell, hex_color="1A2B68"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell(cell, text, bold=False, center=False, size=11, white=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=RGBColor(255, 255, 255) if white else None)


def add_table(doc, headers, rows, header_color="1A2B68"):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        set_cell(table.rows[0].cells[i], h, bold=True, center=True, white=True)
        shade_cell(table.rows[0].cells[i], header_color)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            set_cell(cells[i], val, size=11)
    doc.add_paragraph()
    return table


def add_caption(doc, text):
    add_para(doc, text, size=10, italic=True, center=True, first_indent=False, space_after=12, color=RGBColor(0x47, 0x55, 0x69))


def add_picture(doc, path, width=6.2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))


def build():
    ensure_assets()
    arch = make_architecture_png()
    sitemap = make_sitemap_png()
    usecase = make_usecase_png()
    flow = make_flow_png()

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    # Cover
    for line in [
        "Kingdom of Saudi Arabia",
        "Ministry of Education",
        "University of Hail",
        "College of Computer Science and Engineering",
        "Cooperative Training Office",
    ]:
        add_para(doc, line, size=13, bold=True, center=True, space_after=2, first_indent=False)

    add_para(doc, "", center=True, first_indent=False)
    add_para(doc, "FINAL COOPERATIVE TRAINING REPORT", size=18, bold=True, center=True, first_indent=False, color=RGBColor(0x1A, 0x2B, 0x68))
    add_para(doc, "Project Title", size=12, center=True, first_indent=False, space_after=2)
    add_para(doc, "Nadee Smart Assistant", size=16, bold=True, center=True, first_indent=False, color=RGBColor(0x1A, 0x2B, 0x68))
    add_para(
        doc,
        "An Internal Intelligent Support System for the Data Management Office\nat Hail Region Municipality",
        size=12, center=True, first_indent=False, space_after=16,
    )

    info = [
        ("Trainee Name", "Elaf Abdulaziz Alkowair"),
        ("Training Organization", "Hail Region Municipality"),
        ("Training Department", "Data Management Office"),
        ("Faculty Advisor", "Omar Salem AL-Mostah"),
        ("Academic College", "College of Computer Science and Engineering"),
        ("Report Type", "Final Academic Coop Report"),
        ("Academic Year", "1446–1447 AH / 2025–2026"),
    ]
    add_table(doc, ["Field", "Details"], info)

    add_para(doc, "Submitted in partial fulfillment of the Cooperative Training Program requirements.", size=11, italic=True, center=True, first_indent=False)
    doc.add_page_break()

    # Abstract
    add_h(doc, "Abstract", 1)
    add_para(
        doc,
        "This final report presents the cooperative training experience of Elaf Abdulaziz Alkowair at the "
        "Data Management Office of Hail Region Municipality, and documents the applied project "
        "“Nadee Smart Assistant.” The project addresses a practical organizational challenge: employees "
        "frequently spend time searching for information related to the National Data Index (Nadee) and "
        "repeatedly ask similar questions about guides, procedures, maturity criteria, and supporting "
        "documents. The proposed solution is an internal web-based intelligent assistant that answers "
        "Nadee-related questions from a trusted knowledge base derived from SDAIA guidance, and provides "
        "an analytics module to monitor data-management maturity across fourteen domains. The system was "
        "implemented using Python, Django, and SQLite, with an Arabic RTL interface and a GitHub-based "
        "project repository for academic demonstration."
    )
    add_para(doc, "Keywords: Cooperative Training, National Data Index (Nadee), Chatbot, Data Governance, Django, Knowledge Base, Maturity Assessment.", size=11, italic=True, first_indent=False)
    doc.add_page_break()

    # TOC note
    add_h(doc, "Table of Contents", 1)
    toc = [
        "1. Introduction",
        "2. Training Organization Overview",
        "3. Main Tasks and Responsibilities",
        "4. Problem Statement",
        "5. Project Idea",
        "6. Proposed Solution",
        "7. Technologies and Programming Languages",
        "8. System Maps and Diagrams",
        "9. Use Cases",
        "10. System Features and Functional Description",
        "11. Results and Benefits",
        "12. Challenges and Lessons Learned",
        "13. Conclusion and Recommendations",
        "14. References",
        "Appendices",
    ]
    for item in toc:
        add_para(doc, item, first_indent=False, space_after=3)
    add_para(
        doc,
        "Tip: In Microsoft Word, click References → Table of Contents to insert an automatic TOC based on Heading styles, then update page numbers.",
        size=10, italic=True, first_indent=False, color=RGBColor(0x64, 0x74, 0x8B),
    )
    doc.add_page_break()

    # 1 Introduction
    add_h(doc, "1. Introduction", 1)
    add_para(
        doc,
        "Cooperative training is a critical bridge between academic learning and professional practice. "
        "During the internship at Hail Region Municipality, the trainee was immersed in real data-management "
        "operations, including coordination with the IT Department, participation in inter-agency meetings, "
        "and support for reporting activities related to the National Data Index (Nadee)."
    )
    add_para(
        doc,
        "This final report consolidates the training experience—especially insights reflected in monthly "
        "reports five and six—and presents the applied project developed to respond to an operational need "
        "inside the Data Management Office. The report follows an academic structure covering the "
        "organization context, tasks, problem, idea, solution, technologies, maps, and use cases."
    )

    # 2 Organization
    add_h(doc, "2. Training Organization Overview", 1)
    add_h(doc, "2.1 Host Organization", 2)
    add_para(
        doc,
        "The training was conducted at Hail Region Municipality, within the Data Management Office. "
        "The office contributes to data governance, data quality, institutional reporting, and alignment "
        "with national data controls and measurement frameworks under SDAIA."
    )
    add_h(doc, "2.2 Organizational Context", 2)
    add_para(
        doc,
        "The workplace environment requires continuous coordination between business/data teams and "
        "technical teams. During the internship, the trainee supported communication between the Data "
        "Management Office and the IT Department, attended meetings with digital transformation and "
        "statistics stakeholders, and contributed to preparing materials for discussions involving SDAIA. "
        "These activities highlighted both the importance of accurate information and the cost of fragmented "
        "knowledge sources."
    )

    # 3 Tasks
    add_h(doc, "3. Main Tasks and Responsibilities", 1)
    add_para(doc, "Key responsibilities performed during the cooperative training period include:", first_indent=False)
    add_bullet(doc, "Acting as a communication focal point between the Data Management Office and the IT Department.")
    add_bullet(doc, "Participating in meetings with external entities (e.g., Digital Transformation Agency and Statistics Agency), documenting feedback, and preparing required materials.")
    add_bullet(doc, "Supporting Data Quality activities: reviewing, updating, and following up on requirements and reports.")
    add_bullet(doc, "Supporting Data Governance activities by reviewing related files and platform outputs.")
    add_bullet(doc, "Contributing to Nadee-related reporting and organizational data-management follow-up.")
    add_bullet(doc, "Mentoring new trainees by explaining workflows, communication practices, and report preparation.")
    add_bullet(doc, "Self-development through Python learning and professional programs (e.g., IBM AI ELL and Data Science pathways).")

    # 4 Problem
    add_h(doc, "4. Problem Statement", 1)
    add_para(
        doc,
        "Employees in the Data Management Office frequently need reliable answers about the National Data "
        "Index (Nadee), including domains, criteria, maturity levels, supporting documents, and training "
        "materials. In practice, information is distributed across documents and discussions, which leads to:"
    )
    add_bullet(doc, "Difficulty accessing official Nadee guides and procedures quickly.")
    add_bullet(doc, "Repeated similar inquiries that consume staff time.")
    add_bullet(doc, "Delayed preparation for measurement cycles due to fragmented references.")
    add_bullet(doc, "Limited visibility of maturity performance across data-management domains.")
    add_para(
        doc,
        "Therefore, the core problem is the absence of a unified, trusted, and easy-to-use internal "
        "assistant that can provide accurate Nadee knowledge on demand and support maturity follow-up."
    )

    # 5 Idea
    add_h(doc, "5. Project Idea", 1)
    add_para(
        doc,
        "The project idea is to build an internal platform named “Nadee Smart Assistant” dedicated to "
        "municipality employees in the Data Management Office. The platform provides:"
    )
    add_bullet(doc, "An intelligent chatbot that answers Nadee-related questions from a trusted knowledge base.")
    add_bullet(doc, "A structured knowledge section for guides, procedures, FAQs, and training content.")
    add_bullet(doc, "An analytics section that tracks maturity across the fourteen Nadee domains and generates practical alerts.")
    add_para(
        doc,
        "The idea directly links internship observations to a practical software solution that improves "
        "information access, reduces repeated questions, and supports readiness for Nadee measurement cycles."
    )

    # 6 Solution
    add_h(doc, "6. Proposed Solution", 1)
    add_h(doc, "6.1 Solution Overview", 2)
    add_para(
        doc,
        "The implemented solution is a Django-based internal web system with authentication, chatbot "
        "interaction, knowledge browsing, and analytics dashboards. The chatbot uses a lightweight "
        "Arabic-aware matching engine to retrieve the most relevant knowledge article and returns the "
        "answer with its official source reference."
    )
    add_h(doc, "6.2 Core Modules", 2)
    add_table(
        doc,
        ["Module", "Purpose", "Main Output"],
        [
            ("Accounts", "Secure employee access", "Login / logout / registration"),
            ("Chatbot", "Answer Nadee questions", "Conversational responses with sources"),
            ("Knowledge Base", "Organize official content", "Guides, FAQs, procedures, training"),
            ("Analytics", "Monitor maturity performance", "Domain scores, alerts, forecasts"),
            ("Admin Panel", "Manage content and metrics", "CRUD operations for system data"),
        ],
    )
    add_h(doc, "6.3 Knowledge Scope", 2)
    add_para(
        doc,
        "The knowledge base covers the three Nadee measurement elements (practice maturity, compliance, "
        "and operational excellence), fourteen data-management domains totaling forty-two criteria, six "
        "maturity levels, supporting documents (attachments), and readiness guidance for the third "
        "measurement cycle."
    )

    # 7 Technologies
    add_h(doc, "7. Technologies and Programming Languages", 1)
    add_h(doc, "7.1 Programming Languages", 2)
    add_table(
        doc,
        ["Language", "Usage in the Project"],
        [
            ("Python", "Backend logic, chatbot engine, analytics services, management commands"),
            ("HTML", "Page structure and semantic content layout"),
            ("CSS", "Visual design, Nadee-inspired theme, responsive RTL interface"),
            ("JavaScript", "Chat interaction and front-end behavior"),
            ("SQL (via ORM)", "Data persistence through Django models on SQLite"),
        ],
    )
    add_h(doc, "7.2 Frameworks, Tools, and Platforms", 2)
    add_table(
        doc,
        ["Technology", "Role", "Justification"],
        [
            ("Django", "Web framework", "Structured MVC-like architecture, security, admin panel"),
            ("SQLite", "Database", "Lightweight and suitable for internal/academic deployment"),
            ("Git / GitHub", "Version control & demo hosting", "Collaboration, documentation, Pages preview"),
            ("SDAIA Nadee Guide", "Knowledge source", "Official and trusted content foundation"),
            ("WhiteNoise / Gunicorn", "Production readiness options", "Static serving and WSGI deployment support"),
        ],
    )
    add_h(doc, "7.3 Workplace-Related Technologies and Learning", 2)
    add_para(
        doc,
        "In addition to the project stack, the training exposed the trainee to organizational data platforms "
        "and digital initiatives (including Beam-related awareness sessions mentioned in monthly report five), "
        "Nadee reporting practices, and self-development programs in AI and Data Science. These experiences "
        "informed the problem definition and the design of the assistant."
    )

    # 8 Maps
    add_h(doc, "8. System Maps and Diagrams", 1)
    add_h(doc, "8.1 Architecture Diagram", 2)
    add_para(doc, "Figure 1 illustrates the logical architecture of the system layers and data flow.", first_indent=False)
    add_picture(doc, arch, 6.3)
    add_caption(doc, "Figure 1. System architecture of Nadee Smart Assistant.")

    add_h(doc, "8.2 Sitemap", 2)
    add_para(doc, "Figure 2 presents the main navigation structure of the application.", first_indent=False)
    add_picture(doc, sitemap, 6.3)
    add_caption(doc, "Figure 2. Application sitemap.")

    add_h(doc, "8.3 Chatbot Processing Flow", 2)
    add_para(doc, "Figure 3 summarizes how a user question is transformed into a sourced answer.", first_indent=False)
    add_picture(doc, flow, 6.3)
    add_caption(doc, "Figure 3. Chatbot answer-processing flow.")

    add_h(doc, "8.4 Conceptual Data Map", 2)
    add_para(doc, "The main data entities and relationships can be summarized as follows:", first_indent=False)
    add_bullet(doc, "User → owns → ChatSession → contains → ChatMessage")
    add_bullet(doc, "KnowledgeArticle ← referenced by → ChatMessage (optional)")
    add_bullet(doc, "DataDomain stores maturity score, pillar, and criteria count")
    add_bullet(doc, "PerformanceRecord / MaturityFactor / PerformanceAlert support analytics")

    # 9 Use cases
    add_h(doc, "9. Use Cases", 1)
    add_h(doc, "9.1 Use Case Diagram", 2)
    add_para(doc, "Figure 4 shows the primary actors and system use cases.", first_indent=False)
    add_picture(doc, usecase, 6.3)
    add_caption(doc, "Figure 4. Conceptual use-case diagram.")

    add_h(doc, "9.2 Use Case Catalog", 2)
    add_table(
        doc,
        ["ID", "Use Case", "Actor", "Description"],
        [
            ("UC-01", "Authenticate", "Employee", "Log in to access internal modules"),
            ("UC-02", "Ask Assistant", "Employee", "Submit a Nadee-related question and receive an answer"),
            ("UC-03", "Browse Knowledge", "Employee", "Search and open guides/procedures/FAQs"),
            ("UC-04", "View Domain Maturity", "Employee / Supervisor", "Review 14-domain maturity analytics"),
            ("UC-05", "Review Alerts", "Employee / Supervisor", "Read warnings and improvement suggestions"),
            ("UC-06", "Manage Knowledge", "Admin", "Create/update knowledge articles"),
            ("UC-07", "Manage Metrics", "Admin", "Update domain scores, factors, and alerts"),
        ],
    )

    add_h(doc, "9.3 Detailed Scenario (UC-02)", 2)
    add_para(doc, "Actor: Data Management Office employee.", first_indent=False)
    add_para(doc, "Goal: Understand maturity requirements for the Data Quality domain.", first_indent=False)
    add_para(doc, "Main flow:", first_indent=False)
    add_bullet(doc, "Employee logs into the system.")
    add_bullet(doc, "Employee opens the Smart Assistant page.")
    add_bullet(doc, "Employee asks: “What is the Data Quality domain in Nadee?”")
    add_bullet(doc, "System normalizes the question and matches it against the knowledge base.")
    add_bullet(doc, "System returns the best article with source reference (SDAIA / Nadee guide).")
    add_bullet(doc, "Employee may open Analytics to compare the domain maturity score with other domains.")

    # 10 Features
    add_h(doc, "10. System Features and Functional Description", 1)
    add_bullet(doc, "Secure internal access for municipality employees.")
    add_bullet(doc, "Interactive chatbot specialized in Nadee content.")
    add_bullet(doc, "Trusted knowledge base with categorized articles and keywords.")
    add_bullet(doc, "Analytics dashboard for fourteen domains and forty-two criteria.")
    add_bullet(doc, "Alerts and practical recommendations for maturity improvement.")
    add_bullet(doc, "Admin interface for continuous content and metric updates.")
    add_bullet(doc, "Academic preview via GitHub Pages plus full Django runtime locally.")

    # 11 Results
    add_h(doc, "11. Results and Benefits", 1)
    add_para(doc, "The project delivers tangible benefits aligned with training-office needs:", first_indent=False)
    add_bullet(doc, "Faster access to official Nadee information.")
    add_bullet(doc, "Reduction of repetitive inquiries through a shared assistant.")
    add_bullet(doc, "Better readiness for measurement cycles via structured knowledge.")
    add_bullet(doc, "Improved visibility of domain maturity and delay factors.")
    add_bullet(doc, "A complete applied outcome that connects coop experience with software engineering practice.")

    # 12 Challenges
    add_h(doc, "12. Challenges and Lessons Learned", 1)
    add_para(
        doc,
        "Major challenges included balancing daily office responsibilities with mentoring new trainees, "
        "handling detailed governance/quality follow-up, and consolidating scattered knowledge into one "
        "usable system. Another constraint was that GitHub Pages cannot host full Django applications; "
        "therefore, a static preview was prepared for demonstration while the complete system remains "
        "available for local execution. Key lessons include the value of clear requirements, documentation "
        "discipline, knowledge transfer, and designing solutions around real operational pain points."
    )

    # 13 Conclusion
    add_h(doc, "13. Conclusion and Recommendations", 1)
    add_para(
        doc,
        "The cooperative training experience at Hail Region Municipality provided practical exposure to "
        "data management, governance, quality, and inter-departmental coordination. The Nadee Smart "
        "Assistant project translates that experience into an applied internal system that answers "
        "trusted questions, organizes knowledge, and supports maturity monitoring. Future enhancements "
        "may include integrating advanced language models with source validation, connecting analytics "
        "to live operational data, and expanding the knowledge base as new Nadee guidance is released."
    )

    # 14 References
    add_h(doc, "14. References", 1)
    refs = [
        "Saudi Data and AI Authority (SDAIA). National Data Index (Nadee) — Third Measurement Cycle Guide.",
        "University of Hail, College of Computer Science and Engineering, Cooperative Training Office. Monthly Report Forms (S4).",
        "Alkowair, E. A. Monthly Coop Report No. 5 — Hail Region Municipality, Data Management Office.",
        "Alkowair, E. A. Monthly Coop Report No. 6 — Hail Region Municipality, Data Management Office.",
        "Django Software Foundation. Django Documentation. https://docs.djangoproject.com/",
        "Python Software Foundation. Python Documentation. https://docs.python.org/",
        "GitHub. Nadee Smart Assistant Repository. https://github.com/moneerafahaid-collab/nadee-smart-assistant",
    ]
    for i, ref in enumerate(refs, 1):
        add_para(doc, f"[{i}] {ref}", first_indent=False, space_after=4)

    # Appendices
    add_h(doc, "Appendices", 1)
    add_h(doc, "Appendix A: Demo Credentials", 2)
    add_bullet(doc, "Employee: employee / Nadee@123")
    add_bullet(doc, "Admin: admin / Admin@123")

    add_h(doc, "Appendix B: Project Links", 2)
    add_bullet(doc, "Repository: https://github.com/moneerafahaid-collab/nadee-smart-assistant")
    add_bullet(doc, "GitHub Pages preview: https://moneerafahaid-collab.github.io/nadee-smart-assistant/")

    add_h(doc, "Appendix C: Nadee Domains and Criteria Count", 2)
    add_table(
        doc,
        ["#", "Domain", "Criteria"],
        [
            ("1", "Data Governance", "4"),
            ("2", "Metadata and Data Directory", "3"),
            ("3", "Data Quality", "4"),
            ("4", "Data Storage", "3"),
            ("5", "Content and Document Management", "3"),
            ("6", "Data Modeling and Architecture", "2"),
            ("7", "Reference and Master Data Management", "3"),
            ("8", "Business Intelligence and Analytics", "4"),
            ("9", "Data Integration and Sharing", "4"),
            ("10", "Realizing Value from Data", "2"),
            ("11", "Open Data", "3"),
            ("12", "Freedom of Information", "2"),
            ("13", "Data Classification", "3"),
            ("14", "Personal Data Protection", "2"),
            ("", "Total", "42"),
        ],
    )

    add_h(doc, "Appendix D: Report Content Distribution", 2)
    add_table(
        doc,
        ["Section Group", "Content Focus", "Approx. Weight"],
        [
            ("Introduction & Organization", "Training context and workplace overview", "15%"),
            ("Tasks & Problem", "Responsibilities and operational pain points", "15%"),
            ("Idea & Solution", "Concept and system design", "25%"),
            ("Technologies & Maps", "Languages, tools, diagrams", "20%"),
            ("Use Cases & Features", "Functional modeling", "15%"),
            ("Results & Conclusion", "Outcomes, lessons, recommendations", "10%"),
        ],
    )

    add_para(doc, "", first_indent=False, space_after=16)
    add_para(doc, "Prepared by:", center=True, first_indent=False, space_after=2)
    add_para(doc, "Elaf Abdulaziz Alkowair", bold=True, center=True, first_indent=False, space_after=2)
    add_para(doc, "Data Management Office — Hail Region Municipality", center=True, first_indent=False)

    doc.save(OUT)
    print("Saved:", OUT)


if __name__ == "__main__":
    build()
